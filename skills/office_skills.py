import os
import sys
from docx import Document

def esegui_crea_preventivo_word(dati):
    """Genera un file Word formattato per un preventivo sul Desktop."""
    if not dati:
        return "Impossibile generare il preventivo: dati mancanti."
    
    # Estrai i dati dal dizionario inviato dall'LLM
    mittente = dati.get("azienda_mittente", "N/D")
    ricevente = dati.get("azienda_ricevente", "N/D")
    valore = dati.get("valore", "0")
    iva = dati.get("iva", "22")
    
    # Calcolo totale prudenziale
    try:
        valore_num = float(valore.replace("€", "").replace(",", ".").strip())
        iva_num = float(iva.replace("%", "").strip())
        totale = valore_num + (valore_num * (iva_num / 100))
        totale_str = f"{totale:.2f} €"
    except Exception:
        totale_str = "Da calcolare (valori non numerici)"

    # Crea il documento Word reale
    doc = Document()
    doc.add_heading(f"PREVENTIVO COMMERCIALIZZAZIONE PRODOTTI", level=0)
    
    doc.add_paragraph(f"Società Mittente: {mittente}")
    doc.add_paragraph(f"Cliente Spett.le: {ricevente}")
    doc.add_paragraph("-" * 40)
    
    doc.add_paragraph(f"Valore Merce (Imponibile): {valore}")
    doc.add_paragraph(f"Aliquota IVA applicata: {iva}%")
    doc.add_heading(f"TOTALE PREVENTIVATO (IVA Inclusa): {totale_str}", level=2)
    
    # Percorso Desktop di Windows o Linux
    home = os.path.expanduser("~")
    percorso_desktop = os.path.join(home, "Desktop")
    nome_file = f"Preventivo_{ricevente.replace(' ', '_')}.docx"
    percorso_finale = os.path.join(percorso_desktop, nome_file)
    
    try:
        doc.save(percorso_finale)
        return f"File Word '{nome_file}' generato con successo sul tuo Desktop."
    except Exception as e:
        return f"Errore durante il salvataggio del file Word: {e}"
